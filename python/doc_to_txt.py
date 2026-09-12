#!/usr/bin/env python3
"""
docviewer.py — a multi-format "document -> styled HTML" compiler.

This is the Python sibling of your browser-based Universal Document
Viewer: same idea (a FORMAT_REGISTRY of {extensions, render} entries,
same "A4 page" HTML/CSS shell), but it runs from the command line and
writes a single self-contained .html file instead of rendering live in
a <div>.

Unlike the first docx-only version of this tool, this one leans on a
handful of well-chosen libraries instead of staying lib-less — the
same way your viewer pulls in pdf.js, JSZip, SheetJS, RTF.js, etc.
Reimplementing a PDF renderer or an XLSX formula engine from scratch
isn't a good use of anyone's time; the libraries below are the
standard, well-maintained choice for each format:

    pip install pymupdf openpyxl striprtf markdown --break-system-packages

Everything else (docx, odt, csv, txt, json, images, html) is handled
with just the standard library, the same way the original docx2html.py
was, since those formats don't need much more than a zip/XML/CSV reader.

Usage:
    python3 docviewer.py input.ANYFORMAT [output.html]

To add a new format: write one `render(path) -> RenderResult` function
and append one entry to FORMAT_REGISTRY. Nothing else needs to change
— exactly like the comment in your JS file's FORMAT_REGISTRY.
"""
import sys
import os
import csv
import io
import json
import zipfile
import base64
import html
import xml.etree.ElementTree as ET
from dataclasses import dataclass, field

# ---------------------------------------------------------------------------
# Shared "A4 page" HTML shell (same look as the original docx2html.py output)
# ---------------------------------------------------------------------------

PAGE_TEMPLATE = """<!DOCTYPE html>
<html>
<head>
<meta charset="utf-8">
<title>{title}</title>
<style>
  html, body {{
    margin: 0; padding: 0; background: #d9d9d9;
    font-family: "Calibri", "Segoe UI", Arial, sans-serif;
    font-size: 11pt; line-height: 1.4;
  }}
  .page-wrapper {{ display: flex; flex-direction: column; align-items: center;
    padding: 32px 16px; box-sizing: border-box; min-height: 100vh; gap: 24px; }}
  .page {{ width: 794px; max-width: 100%; min-height: 1123px; background: #fff;
    box-shadow: 0 0 12px rgba(0,0,0,.25); padding: 96px; box-sizing: border-box;
    overflow-wrap: break-word; }}
  .page p {{ margin: 0 0 8pt 0; }}
  .page table {{ width: 100%; margin-bottom: 8pt; border-collapse: collapse; }}
  .page td, .page th {{ padding: 4px 8px; vertical-align: top; border: 1px solid #d8dce4; }}
  .page th {{ background: #f3f4f7; font-weight: 600; }}
  .page ul, .page ol {{ margin: 0 0 8pt 0; padding-left: 24px; }}
  .page img {{ max-width: 100%; }}
  .page pre {{ white-space: pre-wrap; word-break: break-word; font-family: "Consolas","Menlo",monospace; font-size: 10pt; margin: 0; }}
  .pdf-page {{ padding: 0; }}
  .pdf-page img {{ display: block; width: 100%; }}
  .sheet-tabs {{ display: flex; gap: 4px; margin-bottom: 12px; flex-wrap: wrap; }}
  .sheet-tabs button {{ font: inherit; padding: 6px 12px; border: 1px solid #cfd4dd;
    background: #f3f4f7; border-radius: 6px 6px 0 0; cursor: pointer; }}
  .sheet-tabs button.active {{ background: #fff; border-bottom: 2px solid #4c8bf5; font-weight: 600; }}
  .sheet {{ display: none; }}
  .sheet.active {{ display: block; }}
  .note {{ font-family: system-ui, sans-serif; font-size: 12px; color: #888; margin-top: 12px; }}
</style>
</head>
<body>
  <div class="page-wrapper">
{body}
  </div>
<script>
  // Minimal spreadsheet-style tab switcher, mirrors the xlsx tab UI in the
  // browser viewer this was ported from.
  document.querySelectorAll(".sheet-tabs").forEach(function (tabs) {{
    tabs.addEventListener("click", function (e) {{
      var btn = e.target.closest("button");
      if (!btn) return;
      var wrap = tabs.parentElement;
      tabs.querySelectorAll("button").forEach(function (b) {{ b.classList.remove("active"); }});
      wrap.querySelectorAll(".sheet").forEach(function (s) {{ s.classList.remove("active"); }});
      btn.classList.add("active");
      wrap.querySelector('[data-sheet="' + btn.dataset.idx + '"]').classList.add("active");
    }});
  }});
</script>
</body>
</html>
"""


@dataclass
class RenderResult:
    title: str
    body_html: str      # one or more <div class="page">...</div> blocks
    status: str = ""
    note: str = ""       # shown as a small disclaimer under the content, if any


# ---------------------------------------------------------------------------
# DOCX  (stdlib only — zipfile + ElementTree; ported from docx2html.py)
# ---------------------------------------------------------------------------

W_NS = "http://schemas.openxmlformats.org/wordprocessingml/2006/main"
R_NS = "http://schemas.openxmlformats.org/officeDocument/2006/relationships"
A_NS = "http://schemas.openxmlformats.org/drawingml/2006/main"
WP_NS = "http://schemas.openxmlformats.org/drawingml/2006/wordprocessingDrawing"


def _wqn(tag):
    return f"{{{W_NS}}}{tag}"


def _wpqn(tag):
    return f"{{{WP_NS}}}{tag}"


def _emu_to_px(val):
    try:
        return round(int(val) / 9525)  # 914400 EMU = 96px
    except (TypeError, ValueError):
        return None


HIGHLIGHTS = {
    "yellow": "#FFFF00", "green": "#00FF00", "cyan": "#00FFFF", "magenta": "#FF00FF",
    "blue": "#0000FF", "red": "#FF0000", "darkBlue": "#00008B", "darkCyan": "#008B8B",
    "darkGreen": "#006400", "darkMagenta": "#8B008B", "darkRed": "#8B0000",
    "darkYellow": "#808000", "darkGray": "#A9A9A9", "lightGray": "#D3D3D3",
    "black": "#000000", "white": "#FFFFFF", "none": None,
}

# Word list formats that should render as an ordered list (<ol> + CSS
# list-style-type); everything else (bullet, none, etc.) falls back to <ul>.
ORDERED_FORMATS = {
    "decimal": "decimal", "decimalZero": "decimal",
    "lowerLetter": "lower-alpha", "upperLetter": "upper-alpha",
    "lowerRoman": "lower-roman", "upperRoman": "upper-roman",
}


class _DocxRun:
    __slots__ = ("text", "bold", "italic", "underline", "strike", "color",
                 "highlight", "size", "image_html", "is_break", "is_tab")

    def __init__(self, text="", bold=False, italic=False, underline=False,
                 strike=False, color=None, highlight=None, size=None,
                 image_html=None, is_break=False, is_tab=False):
        self.text, self.bold, self.italic = text, bold, italic
        self.underline, self.strike = underline, strike
        self.color, self.highlight, self.size = color, highlight, size
        self.image_html, self.is_break, self.is_tab = image_html, is_break, is_tab


class _DocxHyperlink:
    """A run of hyperlinked text; wraps the runs inside a w:hyperlink so the
    href can be rendered as a real <a> tag instead of being dropped."""
    __slots__ = ("href", "runs")

    def __init__(self, href, runs):
        self.href = href
        self.runs = runs


class _DocxParagraph:
    def __init__(self, runs=None, align=None, indent_left=0, indent_right=0,
                 list_id=None, list_ilvl=0, list_ordered=False, list_style="disc"):
        self.runs = runs or []  # items: _DocxRun or _DocxHyperlink
        self.align, self.indent_left, self.indent_right = align, indent_left, indent_right
        self.list_id, self.list_ilvl = list_id, list_ilvl
        self.list_ordered, self.list_style = list_ordered, list_style


class _DocxTable:
    def __init__(self, rows):
        self.rows = rows  # list[list[list[block]]]


def _render_docx_run(r):
    if r.is_break:
        return "<br>"
    if r.is_tab:
        return "&#9;"
    if r.image_html:
        return r.image_html
    parts = []
    if r.highlight:
        parts.append(f"background-color:{r.highlight}")
    parts.append(f"color:{r.color or '#000000'}")
    if r.size:
        parts.append(f"font-size:{r.size:.1f}pt")
    if r.bold:
        parts.append("font-weight:bold")
    if r.italic:
        parts.append("font-style:italic")
    if r.underline:
        parts.append("text-decoration:underline")
    if r.strike:
        parts.append("text-decoration:line-through")
    text = html.escape(r.text)
    if not text:
        return ""
    return f'<span style="{";".join(parts)}">{text}</span>'


def _render_docx_item(item):
    if isinstance(item, _DocxHyperlink):
        inner = "".join(_render_docx_run(r) for r in item.runs)
        return f'<a href="{html.escape(item.href)}">{inner}</a>'
    return _render_docx_run(item)


def _render_docx_paragraph(p):
    parts = []
    if p.align:
        parts.append(f"text-align:{p.align}")
    if p.indent_left:
        parts.append(f"margin-left:{p.indent_left}px")
    if p.indent_right:
        parts.append(f"margin-right:{p.indent_right}px")
    style = f' style="{";".join(parts)}"' if parts else ""
    return f"<p{style}>{''.join(_render_docx_item(r) for r in p.runs)}</p>"


def _render_docx_blocks(blocks):
    """Renders a sequence of paragraphs/tables to HTML, opening/closing
    nested <ol>/<ul> as the list level (ilvl) rises and falls across
    consecutive list paragraphs, and re-opening a fresh list when the
    numId changes at the same depth (e.g. one list ends and another
    begins back-to-back)."""
    out = []
    list_stack = []  # list of (num_id, ilvl, ordered)

    def close_lists_to(level):
        while len(list_stack) > level:
            _, _, ordered = list_stack.pop()
            out.append("</ol>" if ordered else "</ul>")

    for block in blocks:
        if isinstance(block, _DocxTable):
            close_lists_to(0)
            out.append('<table border="1" style="border-collapse:collapse">')
            for row_cells in block.rows:
                out.append("<tr>")
                for cell in row_cells:
                    if cell is None or cell.get("skip"):
                        continue
                    attrs = ""
                    if cell["colspan"] > 1:
                        attrs += f' colspan="{cell["colspan"]}"'
                    if cell["rowspan"] > 1:
                        attrs += f' rowspan="{cell["rowspan"]}"'
                    out.append(f"<td{attrs}>{_render_docx_blocks(cell['blocks'])}</td>")
                out.append("</tr>")
            out.append("</table>")
            continue

        if block.list_id is not None:
            target_level = block.list_ilvl + 1
            while len(list_stack) > target_level:
                _, _, ordered = list_stack.pop()
                out.append("</ol>" if ordered else "</ul>")
            if len(list_stack) == target_level and list_stack and (
                    list_stack[-1][0] != block.list_id):
                # Same depth but a different list -> close it and restart.
                _, _, ordered = list_stack.pop()
                out.append("</ol>" if ordered else "</ul>")
            while len(list_stack) < target_level:
                ordered = block.list_ordered if len(list_stack) == target_level - 1 else True
                style = f' style="list-style-type:{block.list_style}"' if ordered else ""
                tag_name = "ol" if ordered else "ul"
                out.append(f"<{tag_name}{style}>")
                list_stack.append((block.list_id, block.list_ilvl, ordered))

            out.append(f"<li>{''.join(_render_docx_item(r) for r in block.runs)}</li>")
            continue

        close_lists_to(0)
        out.append(_render_docx_paragraph(block))

    close_lists_to(0)
    return "\n".join(out)


def _document_title(zf, root, fallback):
    """Best-effort title: docProps/core.xml's dc:title, else the first
    heading paragraph's text, else the given fallback (e.g. filename)."""
    if "docProps/core.xml" in zf.namelist():
        try:
            core = ET.fromstring(zf.read("docProps/core.xml"))
            dc_title = core.find("{http://purl.org/dc/elements/1.1/}title")
            if dc_title is not None and dc_title.text and dc_title.text.strip():
                return dc_title.text.strip()
        except ET.ParseError:
            pass

    body = root.find(_wqn("body"))
    if body is not None:
        for p in body.findall(_wqn("p")):
            ppr = p.find(_wqn("pPr"))
            style_el = ppr.find(_wqn("pStyle")) if ppr is not None else None
            if style_el is not None and (style_el.get(_wqn("val")) or "").startswith("Heading"):
                text = "".join(t.text or "" for t in p.iter(_wqn("t"))).strip()
                if text:
                    return text

    return fallback


def render_docx(path):
    zf = zipfile.ZipFile(path)
    rels = {}
    rels_path = "word/_rels/document.xml.rels"
    if rels_path in zf.namelist():
        for rel in ET.fromstring(zf.read(rels_path)):
            target = rel.get("Target")
            if target and not target.startswith(("http://", "https://")):
                target = "word/" + target.lstrip("/")
            rels[rel.get("Id")] = target

    media_cache = {}

    def image_uri(rid):
        target = rels.get(rid)
        if not target or target not in zf.namelist():
            return None
        if rid in media_cache:
            return media_cache[rid]
        data = zf.read(target)
        ext = target.rsplit(".", 1)[-1].lower()
        mime = {"png": "image/png", "jpg": "image/jpeg", "jpeg": "image/jpeg",
                "gif": "image/gif", "bmp": "image/bmp"}.get(ext, "application/octet-stream")
        uri = f"data:{mime};base64,{base64.b64encode(data).decode('ascii')}"
        media_cache[rid] = uri
        return uri

    # numbering.xml: numId -> {ilvl: numFmt}, read per-level (not just the
    # first level) so nested lists pick the right format at each depth.
    numbering = {}
    if "word/numbering.xml" in zf.namelist():
        nroot = ET.fromstring(zf.read("word/numbering.xml"))
        abstract = {}
        for absNum in nroot.findall(_wqn("abstractNum")):
            an_id = absNum.get(_wqn("abstractNumId"))
            levels = {}
            for lvl in absNum.findall(_wqn("lvl")):
                ilvl = lvl.get(_wqn("ilvl"))
                fmt_el = lvl.find(_wqn("numFmt"))
                fmt = fmt_el.get(_wqn("val")) if fmt_el is not None else "decimal"
                levels[ilvl] = fmt
            abstract[an_id] = levels
        for num in nroot.findall(_wqn("num")):
            num_id = num.get(_wqn("numId"))
            abs_ref = num.find(_wqn("abstractNumId"))
            if abs_ref is not None:
                numbering[num_id] = abstract.get(abs_ref.get(_wqn("val")), {})

    def list_format(num_id, ilvl):
        return numbering.get(num_id, {}).get(ilvl, "bullet")

    # styles.xml: styleId -> (numId, ilvl_or_None), resolved through
    # basedOn inheritance, for numbering attached via a paragraph style
    # (e.g. Word's built-in "List Bullet"/"List Number" styles) rather
    # than directly via w:numPr.
    style_num_pr = {}
    if "word/styles.xml" in zf.namelist():
        sroot = ET.fromstring(zf.read("word/styles.xml"))
        raw, based_on = {}, {}
        for style in sroot.findall(_wqn("style")):
            style_id = style.get(_wqn("styleId"))
            based = style.find(_wqn("basedOn"))
            if based is not None:
                based_on[style_id] = based.get(_wqn("val"))
            ppr = style.find(_wqn("pPr"))
            if ppr is None:
                continue
            num_pr = ppr.find(_wqn("numPr"))
            if num_pr is None:
                continue
            num_id_el = num_pr.find(_wqn("numId"))
            ilvl_el = num_pr.find(_wqn("ilvl"))
            if num_id_el is not None:
                raw[style_id] = (num_id_el.get(_wqn("val")),
                                  ilvl_el.get(_wqn("val")) if ilvl_el is not None else None)

        def resolve_style_num(style_id, seen=None):
            if style_id in raw:
                return raw[style_id]
            seen = seen or set()
            parent = based_on.get(style_id)
            if parent and parent not in seen:
                return resolve_style_num(parent, seen | {style_id})
            return None

        style_num_pr = {sid: resolve_style_num(sid) for sid in set(raw) | set(based_on)}

    def parse_rPr(rPr):
        p = dict(bold=False, italic=False, underline=False, strike=False,
                  color=None, highlight=None, size=None)
        if rPr is None:
            return p
        b = rPr.find(_wqn("b"))
        if b is not None and b.get(_wqn("val")) != "false":
            p["bold"] = True
        i = rPr.find(_wqn("i"))
        if i is not None and i.get(_wqn("val")) != "false":
            p["italic"] = True
        u = rPr.find(_wqn("u"))
        if u is not None and u.get(_wqn("val")) not in (None, "none"):
            p["underline"] = True
        if rPr.find(_wqn("strike")) is not None:
            p["strike"] = True
        color = rPr.find(_wqn("color"))
        if color is not None:
            v = color.get(_wqn("val"))
            if v and v != "auto":
                p["color"] = "#" + v.upper()
        hl = rPr.find(_wqn("highlight"))
        if hl is not None:
            p["highlight"] = HIGHLIGHTS.get(hl.get(_wqn("val")))
        sz = rPr.find(_wqn("sz"))
        if sz is not None and sz.get(_wqn("val")):
            p["size"] = int(sz.get(_wqn("val"))) / 2.0
        return p

    def parse_run(r_el):
        props = parse_rPr(r_el.find(_wqn("rPr")))
        runs = []
        for child in r_el:
            if child.tag == _wqn("t"):
                runs.append(_DocxRun(text=child.text or "", **props))
            elif child.tag == _wqn("tab"):
                runs.append(_DocxRun(is_tab=True))
            elif child.tag in (_wqn("br"), _wqn("cr")):
                runs.append(_DocxRun(is_break=True))
            elif child.tag == _wqn("drawing"):
                blip = child.find(f".//{{{A_NS}}}blip")
                if blip is not None:
                    uri = image_uri(blip.get(f"{{{R_NS}}}embed"))
                    if uri:
                        extent = child.find(f".//{_wpqn('extent')}")
                        dims = ""
                        if extent is not None:
                            w_px = _emu_to_px(extent.get("cx"))
                            h_px = _emu_to_px(extent.get("cy"))
                            if w_px and h_px:
                                dims = f' width="{w_px}" height="{h_px}"'
                        runs.append(_DocxRun(image_html=f'<img src="{uri}"{dims} style="max-width:100%">'))
        return runs

    def paragraph_list_info(pPr):
        """Returns {num_id, ilvl, ordered, list_style} or None, resolving
        numbering either from a direct w:numPr or (falling back) from the
        paragraph's style, e.g. Word's built-in "List Bullet"/"List
        Number" styles."""
        if pPr is None:
            return None
        num_id, ilvl = None, "0"

        num_pr = pPr.find(_wqn("numPr"))
        if num_pr is not None:
            ilvl_el = num_pr.find(_wqn("ilvl"))
            num_id_el = num_pr.find(_wqn("numId"))
            if num_id_el is not None:
                num_id = num_id_el.get(_wqn("val"))
                ilvl = ilvl_el.get(_wqn("val")) if ilvl_el is not None else "0"

        if num_id is None:
            style_el = pPr.find(_wqn("pStyle"))
            if style_el is not None:
                resolved = style_num_pr.get(style_el.get(_wqn("val")))
                if resolved:
                    num_id, style_ilvl = resolved
                    if style_ilvl is not None:
                        ilvl = style_ilvl

        if num_id is None or num_id == "0":
            return None
        fmt = list_format(num_id, ilvl)
        is_ordered = fmt in ORDERED_FORMATS
        return {"num_id": num_id, "ilvl": int(ilvl), "ordered": is_ordered,
                "list_style": ORDERED_FORMATS.get(fmt, "disc")}

    def parse_paragraph(p_el):
        pPr = p_el.find(_wqn("pPr"))
        align, indent_left, indent_right = None, 0, 0
        if pPr is not None:
            jc = pPr.find(_wqn("jc"))
            if jc is not None:
                align = {"both": "justify", "start": "left", "end": "right"}.get(jc.get(_wqn("val")), jc.get(_wqn("val")))
            ind = pPr.find(_wqn("ind"))
            if ind is not None:
                left = ind.get(_wqn("left")) or ind.get(_wqn("start"))
                right = ind.get(_wqn("right")) or ind.get(_wqn("end"))
                if left:
                    try:
                        indent_left = round(int(left) / 20 * 96 / 72)
                    except ValueError:
                        pass
                if right:
                    try:
                        indent_right = round(int(right) / 20 * 96 / 72)
                    except ValueError:
                        pass

        list_info = paragraph_list_info(pPr)
        list_id = list_info["num_id"] if list_info else None
        list_ilvl = list_info["ilvl"] if list_info else 0
        list_ordered = list_info["ordered"] if list_info else False
        list_style = list_info["list_style"] if list_info else "disc"

        runs = []
        for child in p_el:
            if child.tag == _wqn("r"):
                runs.extend(parse_run(child))
            elif child.tag == _wqn("hyperlink"):
                r_id = child.get(f"{{{R_NS}}}id")
                href = rels.get(r_id) or "#"
                hyperlink_runs = []
                for r in child.findall(_wqn("r")):
                    hyperlink_runs.extend(parse_run(r))
                runs.append(_DocxHyperlink(href=href, runs=hyperlink_runs))
        return _DocxParagraph(runs=runs, align=align, indent_left=indent_left, indent_right=indent_right,
                               list_id=list_id, list_ilvl=list_ilvl,
                               list_ordered=list_ordered, list_style=list_style)

    def parse_table(tbl_el):
        # grid[row_idx] is a list of cell-descriptor dicts (or None for a
        # slot consumed by a colspan/rowspan from elsewhere). OOXML gives
        # every merged-into cell its own explicit <w:tc> (vMerge="continue")
        # at the same grid column, so column positions never need guessing.
        grid = []
        open_spans = {}  # col_idx -> descriptor currently growing via vMerge

        for row in tbl_el.findall(_wqn("tr")):
            row_cells = []
            col_idx = 0

            def place(idx, descriptor, _row_cells=row_cells):
                while len(_row_cells) <= idx:
                    _row_cells.append(None)
                _row_cells[idx] = descriptor

            for tc in row.findall(_wqn("tc")):
                tcPr = tc.find(_wqn("tcPr"))
                span = 1
                vmerge_val = None
                if tcPr is not None:
                    gs = tcPr.find(_wqn("gridSpan"))
                    if gs is not None:
                        span = int(gs.get(_wqn("val"), "1"))
                    vm = tcPr.find(_wqn("vMerge"))
                    if vm is not None:
                        vmerge_val = vm.get(_wqn("val")) or "continue"

                if vmerge_val == "continue" and col_idx in open_spans:
                    open_spans[col_idx]["rowspan"] += 1
                    place(col_idx, {"skip": True})
                    col_idx += span
                    continue

                descriptor = {"skip": False, "colspan": span, "rowspan": 1,
                               "blocks": parse_blocks(tc)}
                place(col_idx, descriptor)
                if vmerge_val == "restart":
                    open_spans[col_idx] = descriptor
                elif col_idx in open_spans:
                    del open_spans[col_idx]
                col_idx += span

            grid.append(row_cells)
        return _DocxTable(grid)

    def parse_blocks(container):
        blocks = []
        for child in container:
            if child.tag == _wqn("p"):
                blocks.append(parse_paragraph(child))
            elif child.tag == _wqn("tbl"):
                blocks.append(parse_table(child))
        return blocks

    root = ET.fromstring(zf.read("word/document.xml"))
    body = root.find(_wqn("body"))
    blocks = parse_blocks(body)
    body_html = f'<div class="page">{_render_docx_blocks(blocks)}</div>'
    title = _document_title(zf, root, fallback=os.path.basename(path))
    return RenderResult(title=title, body_html=body_html)


# ---------------------------------------------------------------------------
# PDF  (via PyMuPDF — rendering each page to an image is the same strategy
# your browser viewer uses with pdf.js + <canvas>, and it's the only way to
# guarantee visual fidelity without reimplementing a PDF layout engine)
# ---------------------------------------------------------------------------

def render_pdf(path):
    import pymupdf  # fitz
    doc = pymupdf.open(path)
    pages_html = []
    for page in doc:
        pix = page.get_pixmap(dpi=150)
        uri = f"data:image/png;base64,{base64.b64encode(pix.tobytes('png')).decode('ascii')}"
        pages_html.append(f'<div class="page pdf-page"><img src="{uri}"></div>')
    n = len(doc)
    return RenderResult(
        title=os.path.basename(path),
        body_html="\n".join(pages_html),
        status=f"{n} page{'s' if n != 1 else ''}",
        note="PDF pages are rendered as images (like pdf.js/&lt;canvas&gt; in a browser), "
             "so text isn't selectable in this output.",
    )


# ---------------------------------------------------------------------------
# XLSX  (via openpyxl)
# ---------------------------------------------------------------------------

def render_xlsx(path):
    import openpyxl
    wb = openpyxl.load_workbook(path, data_only=True)
    tabs, sheets = [], []
    for idx, name in enumerate(wb.sheetnames):
        ws = wb[name]
        active = " active" if idx == 0 else ""
        tabs.append(f'<button data-idx="{idx}"{" class=\"active\"" if idx == 0 else ""}>{html.escape(name)}</button>')
        rows_html = []
        for row in ws.iter_rows():
            cells = "".join(f"<td>{html.escape(str(c.value)) if c.value is not None else ''}</td>" for c in row)
            rows_html.append(f"<tr>{cells}</tr>")
        sheets.append(f'<div class="sheet{active}" data-sheet="{idx}"><table>{"".join(rows_html)}</table></div>')
    body_html = (
        '<div class="page"><div class="sheet-tabs">' + "".join(tabs) + "</div>"
        + "".join(sheets) + "</div>"
    )
    return RenderResult(title=os.path.basename(path), body_html=body_html,
                         status=f"{len(wb.sheetnames)} sheet(s)")


# ---------------------------------------------------------------------------
# CSV / TSV  (stdlib csv)
# ---------------------------------------------------------------------------

def render_csv(path):
    delim = "\t" if path.lower().endswith((".tsv", ".tab")) else ","
    with open(path, newline="", encoding="utf-8-sig", errors="replace") as f:
        rows = list(csv.reader(f, delimiter=delim))
    if not rows:
        return RenderResult(title=os.path.basename(path), body_html='<div class="page"><p>(empty file)</p></div>')
    header, body = rows[0], rows[1:]
    thead = "".join(f"<th>{html.escape(c)}</th>" for c in header)
    tbody = "".join("<tr>" + "".join(f"<td>{html.escape(c)}</td>" for c in r) + "</tr>" for r in body)
    table = f"<table><thead><tr>{thead}</tr></thead><tbody>{tbody}</tbody></table>"
    return RenderResult(title=os.path.basename(path), body_html=f'<div class="page">{table}</div>',
                         status=f"{len(body)} rows x {len(header)} cols")


# ---------------------------------------------------------------------------
# Plain text / JSON / logs / config  (stdlib)
# ---------------------------------------------------------------------------

def render_text(path):
    with open(path, encoding="utf-8", errors="replace") as f:
        text = f.read()
    return RenderResult(title=os.path.basename(path),
                         body_html=f'<div class="page"><pre>{html.escape(text)}</pre></div>',
                         status=f"{len(text)} chars")


def render_json(path):
    with open(path, encoding="utf-8", errors="replace") as f:
        raw = f.read()
    try:
        pretty = json.dumps(json.loads(raw), indent=2, ensure_ascii=False)
    except json.JSONDecodeError:
        pretty = raw
    return RenderResult(title=os.path.basename(path),
                         body_html=f'<div class="page"><pre>{html.escape(pretty)}</pre></div>')


# ---------------------------------------------------------------------------
# Markdown  (via the `markdown` library)
# ---------------------------------------------------------------------------

def render_markdown(path):
    import markdown as md
    with open(path, encoding="utf-8", errors="replace") as f:
        text = f.read()
    html_body = md.markdown(text, extensions=["tables", "fenced_code"])
    return RenderResult(title=os.path.basename(path), body_html=f'<div class="page">{html_body}</div>')


# ---------------------------------------------------------------------------
# RTF  (via striprtf — text-only; RTF's own formatting model is dropped,
# same trade-off your RTF.js handler would hit for anything beyond basic runs)
# ---------------------------------------------------------------------------

def render_rtf(path):
    from striprtf.striprtf import rtf_to_text
    with open(path, encoding="utf-8", errors="replace") as f:
        raw = f.read()
    text = rtf_to_text(raw)
    paras = "".join(f"<p>{html.escape(p)}</p>" for p in text.split("\n"))
    return RenderResult(title=os.path.basename(path), body_html=f'<div class="page">{paras}</div>',
                         note="RTF is converted to plain text; character-level formatting (bold/italic/color) is not preserved.")


# ---------------------------------------------------------------------------
# ODT  (OpenDocument Text — also a zip of XML, so stdlib zipfile/ElementTree
# is enough for a paragraph-level read, same approach as DOCX above)
# ---------------------------------------------------------------------------

ODT_TEXT_NS = "urn:oasis:names:tc:opendocument:xmlns:text:1.0"


def render_odt(path):
    zf = zipfile.ZipFile(path)
    root = ET.fromstring(zf.read("content.xml"))
    paras = []
    for el in root.iter():
        tag = el.tag.rsplit("}", 1)[-1]
        if tag in ("p", "h"):
            text = "".join(el.itertext())
            paras.append(f"<p>{html.escape(text)}</p>" if text.strip() else "<p></p>")
    return RenderResult(title=os.path.basename(path),
                         body_html=f'<div class="page">{"".join(paras)}</div>',
                         note="ODT is read at paragraph level; character styles (bold/italic/color) are not resolved in this version.")


# ---------------------------------------------------------------------------
# Images  (embed directly as base64)
# ---------------------------------------------------------------------------

IMAGE_MIME = {
    "png": "image/png", "jpg": "image/jpeg", "jpeg": "image/jpeg", "gif": "image/gif",
    "webp": "image/webp", "bmp": "image/bmp", "svg": "image/svg+xml", "ico": "image/x-icon",
}


def render_image(path):
    ext = path.rsplit(".", 1)[-1].lower()
    mime = IMAGE_MIME.get(ext, "application/octet-stream")
    with open(path, "rb") as f:
        data = f.read()
    uri = f"data:{mime};base64,{base64.b64encode(data).decode('ascii')}"
    body = f'<div class="page" style="text-align:center"><img src="{uri}"></div>'
    return RenderResult(title=os.path.basename(path), body_html=body)


# ---------------------------------------------------------------------------
# HTML source view (safe — shown as escaped text, not executed)
# ---------------------------------------------------------------------------

def render_html_source(path):
    with open(path, encoding="utf-8", errors="replace") as f:
        text = f.read()
    return RenderResult(title=os.path.basename(path),
                         body_html=f'<div class="page"><pre>{html.escape(text)}</pre></div>',
                         status="Source view (not executed)")


# ---------------------------------------------------------------------------
# FORMAT REGISTRY — extensions -> render function.
# To add a format: write a `render(path) -> RenderResult` and add one row.
# ---------------------------------------------------------------------------

FORMAT_REGISTRY = [
    {"id": "docx", "label": "Word Document",
     "extensions": ["docx", "docm", "dotx", "dotm"], "render": render_docx},
    {"id": "pdf", "label": "PDF Document",
     "extensions": ["pdf"], "render": render_pdf},
    {"id": "xlsx", "label": "Excel Workbook",
     "extensions": ["xlsx", "xlsm", "xltx", "xltm"], "render": render_xlsx},
    {"id": "csv", "label": "CSV / TSV",
     "extensions": ["csv", "tsv", "tab"], "render": render_csv},
    {"id": "md", "label": "Markdown",
     "extensions": ["md", "markdown", "mdown"], "render": render_markdown},
    {"id": "rtf", "label": "Rich Text Format",
     "extensions": ["rtf"], "render": render_rtf},
    {"id": "odt", "label": "OpenDocument Text",
     "extensions": ["odt", "fodt"], "render": render_odt},
    {"id": "json", "label": "JSON",
     "extensions": ["json", "geojson"], "render": render_json},
    {"id": "txt", "label": "Plain Text",
     "extensions": ["txt", "log", "ini", "cfg", "conf", "yaml", "yml"], "render": render_text},
    {"id": "image", "label": "Image",
     "extensions": ["png", "jpg", "jpeg", "gif", "webp", "bmp", "svg", "ico"], "render": render_image},
    {"id": "html", "label": "HTML Source",
     "extensions": ["html", "htm", "xhtml"], "render": render_html_source},
]

# Formats intentionally NOT implemented: legacy binary .doc / .xls (pre-2007
# OLE formats) and the rest of the ODF family (.ods/.odp/.odg). These need
# either a full OLE Compound File parser or a real layout engine to do
# properly — the honest options are shelling out to LibreOffice
# (`soffice --headless --convert-to html`) or a paid API, not a from-scratch
# reader. If you need these, LibreOffice headless conversion is the
# pragmatic route; ask and I can wire that in as another registry entry.


def detect_handler(path):
    ext = path.rsplit(".", 1)[-1].lower() if "." in path else ""
    for entry in FORMAT_REGISTRY:
        if ext in entry["extensions"]:
            return entry
    return None


def convert(input_path, output_path=None):
    handler = detect_handler(input_path)
    if handler is None:
        supported = ", ".join(sorted({e for h in FORMAT_REGISTRY for e in h["extensions"]}))
        raise ValueError(f"Unsupported file type for {input_path!r}. Supported: {supported}")

    result = handler["render"](input_path)

    note_html = f'<div class="note">{html.escape(result.note)}</div>' if result.note else ""
    full_html = PAGE_TEMPLATE.format(
        title=html.escape(result.title),
        body=result.body_html + note_html,
    )

    if output_path is not None:
        with open(output_path, "w", encoding="utf-8") as f:
            f.write(full_html)

    return output_path, result, full_html


def main():
    if len(sys.argv) < 2:
        print("Usage: python3 docviewer.py input.<ext> [output.html]", file=sys.stderr)
        print("  With an output path: writes the file.", file=sys.stderr)
        print("  Without one: prints the HTML to stdout, e.g.", file=sys.stderr)
        print("    python3 docviewer.py input.docx > out.html", file=sys.stderr)
        print("Supported: " + ", ".join(sorted({e for h in FORMAT_REGISTRY for e in h["extensions"]})), file=sys.stderr)
        sys.exit(1)
    input_path = sys.argv[1]
    output_path = sys.argv[2] if len(sys.argv) > 2 else None
    out_path, result, full_html = convert(input_path, output_path)

    if out_path is None:
        # No destination given: behave like a normal CLI filter and write
        # the HTML to stdout, so it can be redirected (`> out.html`) or
        # piped. Status info goes to stderr so it never lands in the file.
        sys.stdout.write(full_html)
        if result.status:
            print(result.status, file=sys.stderr)
    else:
        print(f"Wrote {out_path}" + (f"  ({result.status})" if result.status else ""))


if __name__ == "__main__":
    main()
